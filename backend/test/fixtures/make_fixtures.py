"""Regenerates the BOQ upload fixtures used by test/boqParse.test.ts (python3 make_fixtures.py)."""
import csv
from openpyxl import Workbook
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

TITLE = ["BILL OF QUANTITIES - 3 BEDROOM BUNGALOW, LEKKI"]
HEADER = ["Item", "Description", "Unit", "Qty", "Rate (N)", "Amount (N)"]
ROWS = [
    ["", "SUBSTRUCTURE", "", "", "", ""],
    ["A", "Excavate foundation trench", "m3", "24", "4,500", "108,000"],
    ["B", "Concrete 1:2:4 in strip foundation", "m3", "12", "85,000", "1,020,000"],
    ["C", "Labour: setting out and blinding", "item", "1", "60,000", "60,000"],
    ["", "Sub-total carried to summary", "", "", "", "1,188,000"],
    ["", "SUPERSTRUCTURE", "", "", "", ""],
    ["D", "225mm sandcrete blocks", "nr", "1800", "650", "1,170,000"],
    ["E", "Bricklayer workmanship to block walls", "m2", "220", "1,500", "330,000"],
    ["F", "Y12 reinforcement bars", "tonnes", "1.5", "1,100,000", "1,650,000"],
    ["", "GRAND TOTAL", "", "", "", "4,338,000"],
]

wb = Workbook(); ws = wb.active; ws.title = "BOQ"
ws.append(TITLE); ws.append([])
ws.append(HEADER)
for r in ROWS:
    out = []
    for i, v in enumerate(r):
        if i >= 3 and v:
            out.append(float(v.replace(",", "")))  # numeric cells, as a QS would type them
        else:
            out.append(v)
    ws.append(out)
wb.save("boq.xlsx")

with open("boq.csv", "w", newline="") as f:
    w = csv.writer(f); w.writerow(TITLE); w.writerow(HEADER); w.writerows(ROWS)

doc = Document(); doc.add_heading(TITLE[0], 1); doc.add_paragraph("Prepared for client review.")
t = doc.add_table(rows=1, cols=len(HEADER))
for i, h in enumerate(HEADER): t.rows[0].cells[i].text = h
for r in ROWS:
    cells = t.add_row().cells
    for i, v in enumerate(r): cells[i].text = v
doc.save("boq.docx")

c = canvas.Canvas("boq.pdf", pagesize=A4)  # reportlab compresses page streams (FlateDecode)
xs = [40, 70, 300, 350, 400, 480]
y = 800
c.setFont("Helvetica-Bold", 11); c.drawString(40, y, TITLE[0]); y -= 30
c.setFont("Helvetica-Bold", 9)
for x, h in zip(xs, HEADER): c.drawString(x, y, h)
y -= 16; c.setFont("Helvetica", 9)
for r in ROWS:
    for x, v in zip(xs, r):
        if v: c.drawString(x, y, v)
    y -= 14
c.save()
print("fixtures written")
